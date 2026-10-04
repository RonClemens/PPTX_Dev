import importlib
import io
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient

FIXTURE = Path(__file__).parent / "fixtures" / "sample.pptx"
CRM_SAMPLE_XLSX = Path(__file__).parent / "fixtures" / "crm_sample_from_user.xlsx"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# sample.pptx's comments: "0-1" (Alex, on slide 2's bullets), "1-1" (Pat's reply
# to it), "0-2" (Alex, on slide 3's table), "0-3" (Alex, slide-level on slide 4).


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("PPTX_DEV_DATA_DIR", str(tmp_path))
    # config/store read env at import time -> reimport the whole app fresh per test
    for mod in list(sys.modules):
        if mod.startswith("app"):
            del sys.modules[mod]
    main = importlib.import_module("app.main")
    return TestClient(main.app)


def upload(client, data: bytes | None = None, name: str = "sample.pptx") -> str:
    data = data if data is not None else FIXTURE.read_bytes()
    res = client.post("/api/documents", files={"file": (name, data, PPTX_MIME)})
    assert res.status_code == 200, res.text
    return res.json()["meta"]["id"]


def get(client, doc_id: str) -> dict:
    return client.get(f"/api/documents/{doc_id}").json()


def shape(body: dict, slide: int, shape_id: str) -> dict:
    return next(s for s in body["document"]["slides"][slide]["shapes"] if s["id"] == shape_id)


# -- documents ---------------------------------------------------------------------


def test_health_reports_start_time_and_persistent_storage_diagnostics(client, tmp_path):
    """/api/health exists to answer "did the server restart and lose its
    data directory" from outside the app -- e.g. after a report of lost
    work, to tell a real bug apart from a platform restart wiping a
    non-persistent data directory."""
    body = client.get("/api/health").json()
    assert body["status"] == "ok"
    assert body["serverStartedAt"]
    assert body["dataDir"] == str(tmp_path)
    assert body["dataDirIsMountPoint"] is False  # tmp_path is a plain directory, not a mount
    assert body["documentCount"] == 0
    upload(client)
    assert client.get("/api/health").json()["documentCount"] == 1


def test_upload_and_get(client):
    doc_id = upload(client)
    body = get(client, doc_id)
    slides = body["document"]["slides"]
    assert len(slides) == 4 and slides[1]["title"] == "Key Achievements"
    assert set(body["document"]["comments"]) == {"0-1", "1-1", "0-2", "0-3"}
    assert [c["heading"] for c in body["chunks"]] == [
        "Q3 Project Status Review", "Key Achievements", "Schedule", "System Architecture",
    ]
    assert body["meta"]["filename"] == "sample.pptx"


def test_upload_rejects_non_pptx_and_corrupt_files(client):
    res = client.post("/api/documents", files={"file": ("notes.docx", b"x", "application/octet-stream")})
    assert res.status_code == 400 and ".pptx" in res.json()["detail"]
    res = client.post("/api/documents", files={"file": ("broken.pptx", b"not a zip", PPTX_MIME)})
    assert res.status_code == 400 and "could not parse" in res.json()["detail"]


def test_document_payload_includes_comment_ref_guesses(client):
    """Wiring check for the Ref #/Para # derivation: Ref # is the slide the
    comment is on, Para # a requirement ID found in the commented shape."""
    refs = get(client, upload(client))["commentRefs"]
    assert refs["0-1"] == {"ref": "Slide 2", "para": "ACA-HRS-0190"}
    assert refs["0-2"] == {"ref": "Slide 3", "para": None}
    assert refs["0-3"]["ref"] == "Slide 4"


def test_list_and_delete_documents(client):
    doc_id = upload(client)
    assert [d["id"] for d in client.get("/api/documents").json()["documents"]] == [doc_id]
    assert client.delete(f"/api/documents/{doc_id}").status_code == 200
    assert client.get(f"/api/documents/{doc_id}").status_code == 404


def test_markdown_endpoint(client):
    md = client.get(f"/api/documents/{upload(client)}/markdown").json()["markdown"]
    assert "## Slide 3: Schedule" in md and "{>>comment:0-1<<}" in md


def test_media_endpoint_serves_the_slide_picture_and_404s_otherwise(client):
    doc_id = upload(client)
    pic = shape(get(client, doc_id), 3, "s3.sh3")
    res = client.get(f"/api/documents/{doc_id}/media/s3/{pic['relId']}")
    assert res.status_code == 200 and res.headers["content-type"] == "image/png"
    assert res.content[:4] == b"\x89PNG"
    assert client.get(f"/api/documents/{doc_id}/media/s3/rId999").status_code == 404
    assert client.get(f"/api/documents/{doc_id}/media/s9/{pic['relId']}").status_code == 404
    assert client.get(f"/api/documents/{doc_id}/media/bogus/{pic['relId']}").status_code == 404


def test_debug_endpoint_is_structural_only(client):
    doc_id = upload(client)
    body = client.get(f"/api/documents/{doc_id}/debug").json()
    assert body["slide_count"] == 4 and body["parsed_comment_count"] == 4
    assert body["conventional_parts_present"]["ppt/commentAuthors.xml (legacy comments)"] is True
    flat = str(body)
    assert "Key Achievements" not in flat and "Alex Reviewer" not in flat  # never any slide/comment text


def test_export_is_a_valid_pptx_that_python_pptx_opens(client):
    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments", json={"slide_id": "s0", "text": "Add a date", "author": "Jane"})
    res = client.get(f"/api/documents/{doc_id}/export")
    assert res.status_code == 200 and res.headers["content-type"] == PPTX_MIME
    assert zipfile.ZipFile(io.BytesIO(res.content)).testzip() is None
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation(io.BytesIO(res.content))
    assert len(prs.slides) == 4
    assert prs.slides[1].shapes.title.text == "Key Achievements"


def test_reset_restores_the_original_upload(client):
    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments", json={"slide_id": "s0", "text": "x", "author": "J"})
    assert len(get(client, doc_id)["document"]["comments"]) == 5
    body = client.post(f"/api/documents/{doc_id}/reset").json()
    assert len(body["document"]["comments"]) == 4


# -- edits -------------------------------------------------------------------------


def _replace(client, doc_id, run, start, end, new_text, author="Jane Reviewer", end_run=None):
    return client.post(f"/api/documents/{doc_id}/edits/replace-range", json={
        "start_run_id": run, "start_offset": start, "end_run_id": end_run or run, "end_offset": end,
        "new_text": new_text, "author": author,
    })


def test_replace_range_edits_in_place_and_leaves_a_resolved_edit_record(client):
    doc_id = upload(client)
    res = _replace(client, doc_id, "s1.sh3.p2.r0", 0, 24, "Cut unit cost by 12%")
    assert res.status_code == 200, res.text
    body = res.json()
    runs = shape(body, 1, "s1.sh3")["paragraphs"][2]["runs"]
    assert runs[0]["text"] == "Cut unit cost by 12% versus the baseline estimate"
    assert runs[0].get("size") == 32.0  # formatting preserved
    notes = [c for c in body["document"]["comments"].values() if c["kind"] == "edit"]
    assert len(notes) == 1 and notes[0]["done"] is True and notes[0]["author"] == "Jane Reviewer"
    assert "Reduced unit cost by 12%" in notes[0]["text"]
    edits = client.get(f"/api/documents/{doc_id}/history").json()["edits"]
    assert edits[-1]["kind"] == "replace_text_range" and edits[-1]["author"] == "Jane Reviewer"


def test_replace_range_across_paragraphs_is_400(client):
    doc_id = upload(client)
    res = _replace(client, doc_id, "s1.sh3.p0.r0", 0, 3, "x", end_run="s1.sh3.p1.r0")
    assert res.status_code == 400 and "single paragraph" in res.json()["detail"]


def test_replace_range_unknown_run_is_400(client):
    doc_id = upload(client)
    assert _replace(client, doc_id, "s1.sh3.p0.r9", 0, 1, "x").status_code == 400
    assert _replace(client, doc_id, "s1.sh99.p0.r0", 0, 1, "x").status_code == 400
    assert _replace(client, doc_id, "garbage", 0, 1, "x").status_code == 400


# -- comments ----------------------------------------------------------------------


def test_create_comment_via_selection_remembers_the_selected_text(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments", json={
        "start_run_id": "s1.sh3.p3.r0", "start_offset": 14, "end_run_id": "s1.sh3.p3.r0", "end_offset": 33,
        "text": "Which customer?", "author": "Jane Reviewer",
    })
    assert res.status_code == 200, res.text
    c = res.json()["document"]["comments"]["2-1"]
    assert c["quote"] == "first prototype to"  # surrounding whitespace is trimmed
    assert c["shapeId"] == "s1.sh3" and c["slideId"] == "s1" and c["author"] == "Jane Reviewer"


def test_create_slide_and_shape_comments(client):
    doc_id = upload(client)
    a = client.post(f"/api/documents/{doc_id}/comments", json={"slide_id": "s0", "text": "Whole slide", "author": "J"})
    b = client.post(f"/api/documents/{doc_id}/comments", json={
        "slide_id": "s2", "shape_id": "s2.sh3", "text": "On the table", "author": "J"})
    comments = b.json()["document"]["comments"]
    assert comments["2-1"]["shapeId"] is None and comments["2-1"]["slideId"] == "s0"
    assert comments["2-2"]["shapeId"] == "s2.sh3"
    assert a.status_code == 200


def test_create_comment_needs_a_target(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments", json={"text": "x", "author": "J"})
    assert res.status_code == 400
    res = client.post(f"/api/documents/{doc_id}/comments", json={
        "slide_id": "s1", "shape_id": "s1.sh999", "text": "x", "author": "J"})
    assert res.status_code == 400


def test_comment_reply_and_resolve(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/0-2/reply", json={"text": "Confirmed.", "author": "Pat"})
    assert res.status_code == 200, res.text
    comments = res.json()["document"]["comments"]
    reply = next(c for c in comments.values() if c["parentId"] == "0-2")
    assert reply["text"] == "Confirmed." and reply["author"] == "Pat" and reply["shapeId"] == "s2.sh3"

    res = client.post(f"/api/documents/{doc_id}/comments/0-2/resolve", json={"done": True, "author": "Pat"})
    assert res.json()["document"]["comments"]["0-2"]["done"] is True
    res = client.post(f"/api/documents/{doc_id}/comments/0-2/resolve", json={"done": False, "author": "Pat"})
    assert res.json()["document"]["comments"]["0-2"]["done"] is False
    assert client.post(f"/api/documents/{doc_id}/comments/9-9/resolve", json={"done": True}).status_code == 400
    assert client.post(f"/api/documents/{doc_id}/comments/9-9/reply", json={"text": "x"}).status_code == 400


def test_reply_and_resolve_are_in_the_exported_pptx_xml(client):
    """Inspect the raw exported bytes directly (independent of this app's own
    parser), so a matching bug in writer and parser can't hide."""
    from lxml import etree

    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments/0-2/reply", json={"text": "Confirmed.", "author": "Pat Author"})
    client.post(f"/api/documents/{doc_id}/comments/0-2/resolve", json={"done": True, "author": "Pat Author"})
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/documents/{doc_id}/export").content))
    root = etree.fromstring(z.read("ppt/comments/comment2.xml"))
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "p15": "http://schemas.microsoft.com/office/powerpoint/2012/main"}
    cms = root.findall("p:cm", ns)
    assert len(cms) == 2
    reply = cms[1]
    assert reply.find("p:text", ns).text == "Confirmed."
    assert reply.find(".//p15:parentCm", ns).get("idx") == "2"
    assert b"urn:pptx-dev:comment-state" in z.read("ppt/comments/comment2.xml")  # resolved flag lives in our private ext


def test_import_reflects_preexisting_threads_from_real_powerpoint_xml(client):
    """The sample's comments/threads were written as raw XML the way PowerPoint
    does, never by this app -- upload must show them as threads."""
    comments = get(client, upload(client))["document"]["comments"]
    assert comments["1-1"]["parentId"] == "0-1"
    assert comments["0-1"]["parentId"] is None


# -- AI ----------------------------------------------------------------------------


class _FakeToolUse:
    type = "tool_use"

    def __init__(self, name, data):
        self.name, self.input = name, data


def _fake_client(captured: dict, name: str, data: dict):
    class FakeMessages:
        def create(self, **kwargs):
            captured.update(kwargs)
            return SimpleNamespace(content=[_FakeToolUse(name, data)])

    return SimpleNamespace(messages=FakeMessages())


def test_adjudicate_without_api_key_returns_502(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/0-1/adjudicate", json={})
    assert res.status_code == 502 and "No API key available" in res.json()["detail"]


def test_adjudicate_comment_wraps_anthropic_sdk_errors():
    """A real API-level failure (bad key, rate limit, dropped connection) must
    surface as our own AdjudicationError, not an unhandled 500."""
    import anthropic
    import httpx

    import app.ai.assistant as assistant_mod

    fake_client = SimpleNamespace(
        messages=SimpleNamespace(
            create=lambda **kwargs: (_ for _ in ()).throw(
                anthropic.APIConnectionError(request=httpx.Request("POST", "https://api.anthropic.com"))
            )
        )
    )
    with patch.object(assistant_mod, "_client", return_value=fake_client):
        with pytest.raises(assistant_mod.AdjudicationError):
            assistant_mod.adjudicate_comment(
                comment_id="0-1", comment_author="Jane", comment_text="?", context_markdown="text"
            )


def test_suggest_adjudications_reports_anthropic_sdk_error_per_comment(client):
    import anthropic
    import httpx

    import app.ai.assistant as assistant_mod

    fake_client = SimpleNamespace(
        messages=SimpleNamespace(
            create=lambda **kwargs: (_ for _ in ()).throw(
                anthropic.APIConnectionError(request=httpx.Request("POST", "https://api.anthropic.com"))
            )
        )
    )
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=fake_client):
        res = client.post(f"/api/documents/{doc_id}/suggest-adjudications")
    assert res.status_code == 200
    s = res.json()["suggestions"]
    assert set(s) == {"0-1", "0-2", "0-3"}  # open, top-level, non-edit comments only
    assert all("error" in v for v in s.values())


def test_adjudicate_sends_full_deck_context_before_the_comment_ask(client):
    """The AI must see the whole deck (when it fits under the size cap), and
    that context must come before the specific comment it's asked about."""
    import app.ai.assistant as assistant_mod

    captured: dict = {}
    fake = _fake_client(captured, "submit_adjudication", {"action": "reply", "reply": "ok", "reasoning": "ok"})
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=fake):
        res = client.post(f"/api/documents/{doc_id}/comments/0-1/adjudicate", json={})
    assert res.status_code == 200, res.text
    prompt = captured["messages"][0]["content"]
    assert "Full presentation context" in prompt and "Slide context" not in prompt
    # slide 3's table is unrelated to the comment on slide 2, so seeing it proves whole-deck context
    idx_table, idx_ask = prompt.find("Qualification Testing"), prompt.find("Reviewer comment to adjudicate")
    assert 0 <= idx_table < idx_ask
    assert "attached to the marked text box" in prompt  # no stored quote -> told it's a whole-shape comment


def test_adjudicate_tells_the_ai_the_exact_selected_text_when_known(client):
    import app.ai.assistant as assistant_mod

    captured: dict = {}
    fake = _fake_client(captured, "submit_adjudication", {"action": "reply", "reply": "ok", "reasoning": "ok"})
    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments", json={
        "start_run_id": "s1.sh3.p2.r0", "start_offset": 0, "end_run_id": "s1.sh3.p2.r0", "end_offset": 24,
        "text": "Which baseline?", "author": "Jane"})
    with patch.object(assistant_mod, "_client", return_value=fake):
        client.post(f"/api/documents/{doc_id}/comments/2-1/adjudicate", json={})
    prompt = captured["messages"][0]["content"]
    assert 'selected this exact text when commenting: "Reduced unit cost by 12%"' in prompt
    assert "{==Reduced unit cost by 12%==}{>>comment:2-1<<}" in prompt


def test_adjudicate_edit_action_edits_text_replies_resolves_and_records(client):
    import app.ai.assistant as assistant_mod
    import app.api.ai as ai_route_mod

    fake_result = assistant_mod.AdjudicationResult(
        action="edit",
        original_text="verified in lab testing",
        replacement_text="verified in lab testing (TR-2025-014)",
        reply="Added the test report number.",
        reasoning="The comment asks for the report number; it is in the reply thread.",
    )
    doc_id = upload(client)
    with patch.object(ai_route_mod, "adjudicate_comment", return_value=fake_result):
        res = client.post(f"/api/documents/{doc_id}/comments/0-1/adjudicate", json={"author": "AI Assistant"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["adjudication"]["action"] == "edit"
    assert body["adjudication"]["original_text"] == "verified in lab testing"

    text = shape(body, 1, "s1.sh3")["paragraphs"][1]["runs"][0]["text"]
    assert text == "Requirement ACA-HRS-0190 verified in lab testing (TR-2025-014)"
    comments = body["document"]["comments"]
    assert comments["0-1"]["done"] is True
    reply = next(c for c in comments.values() if c["parentId"] == "0-1" and c["author"] == "AI Assistant")
    assert reply["text"] == "Added the test report number."
    record = next(c for c in comments.values() if c["kind"] == "edit")
    assert record["done"] and record["author"] == "AI Assistant" and "TR-2025-014" in record["text"]
    assert body["adjudication"]["change_id"] == record["id"]

    history = client.get(f"/api/documents/{doc_id}/history").json()
    assert history["ai_adjudications"][0]["action"] == "edit"


def test_adjudicate_edit_without_findable_original_text_is_400_and_changes_nothing(client):
    import app.ai.assistant as assistant_mod
    import app.api.ai as ai_route_mod

    fake_result = assistant_mod.AdjudicationResult(
        action="edit", original_text="words that are not on the slide", replacement_text="x",
        reply="done", reasoning="r",
    )
    doc_id = upload(client)
    with patch.object(ai_route_mod, "adjudicate_comment", return_value=fake_result):
        res = client.post(f"/api/documents/{doc_id}/comments/0-1/adjudicate", json={})
    assert res.status_code == 400
    body = get(client, doc_id)
    assert body["document"]["comments"]["0-1"]["done"] is False
    assert not any(c["kind"] == "edit" for c in body["document"]["comments"].values())


def test_adjudicate_no_change_action_leaves_comment_open(client):
    import app.ai.assistant as assistant_mod
    import app.api.ai as ai_route_mod

    fake_result = assistant_mod.AdjudicationResult(
        action="no_change", original_text=None, replacement_text=None,
        reply="This needs the test report from the lab.", reasoning="Insufficient info.",
    )
    doc_id = upload(client)
    with patch.object(ai_route_mod, "adjudicate_comment", return_value=fake_result):
        res = client.post(f"/api/documents/{doc_id}/comments/0-2/adjudicate", json={})
    assert res.status_code == 200, res.text
    comments = res.json()["document"]["comments"]
    assert comments["0-2"]["done"] is False
    assert any(c["parentId"] == "0-2" for c in comments.values())


def test_apply_adjudication_edit_with_explicit_original_without_calling_ai(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/0-1/apply-adjudication", json={
        "action": "edit", "original_text": "ahead of schedule", "replacement_text": "two weeks early",
        "reply": "Made it specific.", "reasoning": "r", "author": "Jane Reviewer",
    })
    assert res.status_code == 200, res.text
    body = res.json()
    assert shape(body, 1, "s1.sh3")["paragraphs"][0]["runs"][0]["text"] == (
        "Completed preliminary design review two weeks early")
    assert body["document"]["comments"]["0-1"]["done"] is True


def test_apply_adjudication_edit_needs_original_for_a_whole_shape_comment(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/0-1/apply-adjudication", json={
        "action": "edit", "replacement_text": "x", "reply": "r"})
    assert res.status_code == 400 and "specify which text" in res.json()["detail"]


def test_apply_adjudication_uses_the_stored_selection_when_there_is_one(client):
    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments", json={
        "start_run_id": "s1.sh3.p2.r0", "start_offset": 0, "end_run_id": "s1.sh3.p2.r0", "end_offset": 24,
        "text": "Wording?", "author": "Jane"})
    res = client.post(f"/api/documents/{doc_id}/comments/2-1/apply-adjudication", json={
        "action": "edit", "replacement_text": "Cut unit cost by 12%", "reply": "Reworded."})
    assert res.status_code == 200, res.text
    assert shape(res.json(), 1, "s1.sh3")["paragraphs"][2]["runs"][0]["text"].startswith("Cut unit cost by 12% versus")


def test_apply_adjudication_unknown_comment_returns_404(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/9-9/apply-adjudication", json={"action": "reply", "reply": "x"})
    assert res.status_code == 404


def test_apply_adjudication_never_silently_drops_a_failed_reply(client):
    """If posting the reply fails, the request must fail loudly rather than
    resolve the comment as though it succeeded."""
    import app.api.ai as ai_route_mod

    doc_id = upload(client)
    with patch.object(ai_route_mod.writer, "add_comment_reply", side_effect=ai_route_mod.writer.EditError("boom")):
        res = client.post(f"/api/documents/{doc_id}/comments/0-2/apply-adjudication", json={
            "action": "reply", "reply": "text that must not vanish"})
    assert res.status_code == 400 and "could not add reply" in res.json()["detail"]
    assert get(client, doc_id)["document"]["comments"]["0-2"]["done"] is False


def test_suggest_adjudications_is_read_only_and_skips_resolved_replies_and_edit_records(client):
    import app.ai.assistant as assistant_mod
    import app.api.ai as ai_route_mod

    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/comments/0-3/resolve", json={"done": True})
    _replace_ok = client.post(f"/api/documents/{doc_id}/edits/replace-range", json={
        "start_run_id": "s1.sh3.p3.r0", "start_offset": 0, "end_run_id": "s1.sh3.p3.r0", "end_offset": 9,
        "new_text": "Shipped", "author": "Jane"})
    assert _replace_ok.status_code == 200
    before = client.get(f"/api/documents/{doc_id}/export").content

    result = assistant_mod.AdjudicationResult(
        action="edit", original_text="ahead of schedule", replacement_text="early", reply="ok", reasoning="r")
    with patch.object(ai_route_mod, "adjudicate_comment", return_value=result) as m:
        res = client.post(f"/api/documents/{doc_id}/suggest-adjudications")
    body = res.json()
    assert set(body["suggestions"]) == {"0-1", "0-2"}  # not resolved 0-3, not reply 1-1, not the edit record
    assert body["suggestions"]["0-1"]["original_text"] == "ahead of schedule"
    assert body["usedFullDocumentContext"] is True
    assert m.call_count == 2
    assert client.get(f"/api/documents/{doc_id}/export").content == before


def test_ai_review_adds_comments_on_the_quoted_text_and_skips_unmatched_quotes(client):
    import app.api.ai as ai_route_mod

    def fake_review(*, section_markdown, doc_title=None):
        if "## Slide 2:" in section_markdown:
            return [
                {"quote": "Reduced unit cost by 12%", "comment": "Versus what baseline?"},
                {"quote": "a quote the model hallucinated", "comment": "dropped"},
            ]
        return []

    doc_id = upload(client)
    with patch.object(ai_route_mod, "generate_review_comments", side_effect=fake_review):
        res = client.post(f"/api/documents/{doc_id}/ai-review", json={"author": "AI Reviewer"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["aiReview"]["added"] == 1 and body["aiReview"]["skipped"] == 1
    new = body["document"]["comments"][body["aiReview"]["commentIds"][0]]
    assert new["author"] == "AI Reviewer" and new["quote"] == "Reduced unit cost by 12%"
    assert new["shapeId"] == "s1.sh3" and new["text"] == "Versus what baseline?"


def test_ai_review_propagates_ai_error(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/ai-review", json={})
    assert res.status_code == 502


# -- decisions / CRM export ----------------------------------------------------------


def _decide(client, doc_id, cid="0-1", decision="accept", reason="Done.", ref="", author="Bob"):
    return client.post(f"/api/documents/{doc_id}/comments/{cid}/decision", json={
        "decision": decision, "reason": reason, "ref": ref, "author": author})


def test_save_comment_decision_records_reason_and_appears_in_payload(client):
    doc_id = upload(client)
    res = _decide(client, doc_id, ref="3.2.2")
    assert res.status_code == 200, res.text
    d = res.json()["decisions"]["0-1"]
    assert (d["decision"], d["reason"], d["ref"], d["author"]) == ("accept", "Done.", "3.2.2", "Bob")
    assert res.json()["decision"]["comment_id"] == "0-1"
    assert get(client, doc_id)["decisions"]["0-1"]["decision"] == "accept"


def test_decision_validation_and_overwrite(client):
    doc_id = upload(client)
    assert _decide(client, doc_id, cid="9-9").status_code == 404
    assert _decide(client, doc_id, decision="maybe").status_code == 422
    _decide(client, doc_id, decision="accept")
    _decide(client, doc_id, decision="defer", reason="Later.")
    assert get(client, doc_id)["decisions"]["0-1"]["decision"] == "defer"


def test_update_document_meta_sets_doc_ref_and_rev(client):
    doc_id = upload(client)
    res = client.patch(f"/api/documents/{doc_id}/meta", json={"doc_ref": "UTIC-MK710-A012-0003", "rev": "X1"})
    assert res.json()["meta"]["doc_ref"] == "UTIC-MK710-A012-0003"
    res = client.patch(f"/api/documents/{doc_id}/meta", json={"rev": "X2"})
    assert res.json()["meta"]["doc_ref"] == "UTIC-MK710-A012-0003" and res.json()["meta"]["rev"] == "X2"


def _tsv_rows(client, doc_id):
    res = client.get(f"/api/documents/{doc_id}/comments/export")
    assert res.status_code == 200 and res.headers["content-type"].startswith("text/tab-separated-values")
    lines = res.text.strip().split("\n")
    return [ln.split("\t") for ln in lines]


def test_export_comment_resolution_matrix_tsv(client):
    doc_id = upload(client)
    client.patch(f"/api/documents/{doc_id}/meta", json={"doc_ref": "UTIC-MK710-A012-0003", "rev": "X1"})
    _decide(client, doc_id, decision="info_only", reason="FYI", ref="")
    rows = _tsv_rows(client, doc_id)
    assert rows[0] == [
        "Item #", "Document", "Rev", "Ref #", "Para #", "Comment",
        "Contractor Response (Accept / Reject)", "Contractor Rationale",
        "Status (Open / Closed)", "Date Adjudicated", "Comment ID",
    ]
    assert len(rows) == 4  # header + 3 top-level comments; the reply isn't its own row
    first = rows[1]
    assert first[1:5] == ["UTIC-MK710-A012-0003", "X1", "Slide 2", "ACA-HRS-0190"]
    assert first[6:9] == ["Info Only", "FYI", "Closed"] and first[10] == "0-1"
    assert rows[2][8] == "Open" and rows[3][3] == "Slide 4" and rows[3][4] == "slide3:s3"  # internal fallback ref


def test_export_excludes_edit_records_and_falls_back_to_filename(client):
    doc_id = upload(client)
    client.post(f"/api/documents/{doc_id}/edits/replace-range", json={
        "start_run_id": "s1.sh3.p3.r0", "start_offset": 0, "end_run_id": "s1.sh3.p3.r0", "end_offset": 9,
        "new_text": "Shipped", "author": "Jane"})
    rows = _tsv_rows(client, doc_id)
    assert len(rows) == 4  # the automatic edit record is not a review comment
    assert rows[1][1] == "sample.pptx"


def test_a_manual_ref_beats_the_derived_one_in_the_export(client):
    doc_id = upload(client)
    _decide(client, doc_id, ref="REQ-7")
    assert _tsv_rows(client, doc_id)[1][3] == "REQ-7"


# -- per-comment AI chat --------------------------------------------------------------


def _chat_fake(captured_list, reply="Plain text answer."):
    class FakeMessages:
        def create(self, **kwargs):
            captured_list.append(kwargs)
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=reply)])

    return SimpleNamespace(messages=FakeMessages())


def test_chat_persists_turns_and_returns_reply(client):
    import app.ai.assistant as assistant_mod

    calls: list = []
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=_chat_fake(calls)):
        res = client.post(f"/api/documents/{doc_id}/comments/0-1/chat", json={"message": "Is this a reject?", "author": "Bob"})
    assert res.status_code == 200, res.text
    history = res.json()["chat"]["history"]
    assert [(t["role"], t["content"]) for t in history] == [
        ("user", "Is this a reject?"), ("assistant", "Plain text answer.")]
    assert history[0]["author"] == "Bob"
    assert get(client, doc_id)["chats"]["0-1"] == history
    assert "Presentation:" in calls[0]["messages"][0]["content"]


def test_chat_second_turn_replays_prior_history_before_new_message(client):
    import app.ai.assistant as assistant_mod

    calls: list = []
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=_chat_fake(calls)):
        client.post(f"/api/documents/{doc_id}/comments/0-1/chat", json={"message": "first", "author": "Bob"})
        client.post(f"/api/documents/{doc_id}/comments/0-1/chat", json={"message": "second", "author": "Bob"})
    msgs = calls[1]["messages"]
    assert [m["role"] for m in msgs] == ["user", "assistant", "user", "assistant", "user"]
    assert msgs[1]["content"].startswith("Understood")
    assert [m["content"] for m in msgs[2:]] == ["first", "Plain text answer.", "second"]


def test_chat_system_prompt_forbids_markdown_so_replies_paste_cleanly(client):
    import app.ai.assistant as assistant_mod

    calls: list = []
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=_chat_fake(calls)):
        client.post(f"/api/documents/{doc_id}/comments/0-1/chat", json={"message": "hi", "author": "Bob"})
    system = calls[0]["system"]
    assert "plain prose only" in system and "no Markdown formatting" in system


def test_chat_unknown_comment_returns_404(client):
    doc_id = upload(client)
    assert client.post(f"/api/documents/{doc_id}/comments/9-9/chat", json={"message": "hi"}).status_code == 404


def test_chat_wraps_anthropic_sdk_errors_as_502_but_keeps_user_message(client):
    import anthropic
    import httpx

    import app.ai.assistant as assistant_mod

    fake = SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: (_ for _ in ()).throw(
        anthropic.APIConnectionError(request=httpx.Request("POST", "https://api.anthropic.com")))))
    doc_id = upload(client)
    with patch.object(assistant_mod, "_client", return_value=fake):
        res = client.post(f"/api/documents/{doc_id}/comments/0-1/chat", json={"message": "hello", "author": "Bob"})
    assert res.status_code == 502
    assert [t["content"] for t in get(client, doc_id)["chats"]["0-1"]] == ["hello"]


# -- XLSX round trip --------------------------------------------------------------------


def _build_crm_xlsx(rows: list[list], header: list[str] | None = None) -> bytes:
    """A minimal stand-in for the real Comment Resolution Matrix template:
    row 1 is a banner (like the real template's merged "Government
    Columns"/"Contractor Columns" row) that must NOT be mistaken for the
    header, row 2 is the real header row, row 3+ is data."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Comments"
    ws.append(["Government Columns", "", "", "", "", "", "Contractor Columns", ""])
    ws.append(header or [
        "Item #", "Document", "Rev", "Ref #", "Para #", "Comment",
        "Contractor Response", "Contractor Rationale", "Status",
    ])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload_xlsx(client, doc_id: str, xlsx_bytes: bytes, author: str = "Bob Contractor"):
    return client.post(
        f"/api/documents/{doc_id}/comments/import-decisions",
        files={"file": ("filled_matrix.xlsx", xlsx_bytes, XLSX_MIME)},
        data={"author": author},
    )


Q = "Can we cite the test report number for the ACA-HRS-0190 verification?"


def test_import_xlsx_matches_by_comment_text_and_pushes_reply_and_resolve(client):
    doc_id = upload(client)
    res = _upload_xlsx(client, doc_id, _build_crm_xlsx([
        [1, "sample.pptx", "X1", "Slide 2", "ACA-HRS-0190", Q, "Accept", "Added TR-2025-014.", ""]]))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["xlsxImport"]["applied"] == [{"row": 3, "comment_id": "0-1", "decision": "accept", "created": False}]
    assert body["xlsxImport"]["skipped"] == []
    assert body["decisions"]["0-1"]["reason"] == "Added TR-2025-014."
    comments = body["document"]["comments"]
    assert comments["0-1"]["done"] is True
    reply = next(c for c in comments.values() if c["parentId"] == "0-1" and c["author"] == "Bob Contractor")
    assert reply["text"] == "Accept: Added TR-2025-014."


def test_imported_disposition_is_in_the_exported_pptx_too(client):
    doc_id = upload(client)
    _upload_xlsx(client, doc_id, _build_crm_xlsx([[1, "", "", "", "", Q, "Reject", "Out of scope.", ""]]))
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/documents/{doc_id}/export").content))
    xml = z.read("ppt/comments/comment1.xml").decode()
    assert "Reject: Out of scope." in xml


def test_import_xlsx_matches_by_hidden_comment_id_over_mismatched_text(client):
    doc_id = upload(client)
    xlsx = _build_crm_xlsx(
        [["1", "sample.pptx", "X1", "", "", "some other unrelated text", "Reject", "Out of scope.", "", "0-2"]],
        header=["Item #", "Document", "Rev", "Ref #", "Para #", "Comment",
                "Contractor Response", "Contractor Rationale", "Status", "Comment ID"],
    )
    body = _upload_xlsx(client, doc_id, xlsx).json()
    assert body["xlsxImport"]["applied"] == [{"row": 3, "comment_id": "0-2", "decision": "reject", "created": False}]
    assert body["decisions"]["0-2"]["decision"] == "reject"


def test_import_xlsx_matches_exact_template_header_wording(client):
    doc_id = upload(client)
    xlsx = _build_crm_xlsx(
        [[1, "sample.pptx", "X1", "", "", Q, "accept", "OK", "Closed"]],
        header=["Item #", "Document", "Rev", "Ref #", "Para #", "Comment",
                "Contractor Response\n(Accept / Reject)", "Contractor Rationale", "Status\n(Open / Closed)"],
    )
    body = _upload_xlsx(client, doc_id, xlsx).json()
    assert body["xlsxImport"]["applied"][0]["comment_id"] == "0-1"


def test_import_xlsx_roundtrips_this_apps_own_export(client):
    doc_id = upload(client)
    rows = _tsv_rows(client, doc_id)
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Comments"
    for i, row in enumerate(rows):
        if i > 0 and row[10] in ("0-1", "0-2"):
            row[6], row[7] = "Accept", f"ok {row[10]}"
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    body = _upload_xlsx(client, doc_id, buf.getvalue()).json()
    assert sorted(a["comment_id"] for a in body["xlsxImport"]["applied"]) == ["0-1", "0-2"]
    # The third row (0-3) had no response -> skipped, as a not-yet-reviewed row.
    assert body["xlsxImport"]["skippedSummary"] == [{"reason": 'no recognized response ("")', "count": 1, "rows": [4]}]


def test_import_xlsx_skips_unrecognized_response_and_unmatched_comment(client):
    doc_id = upload(client)
    body = _upload_xlsx(client, doc_id, _build_crm_xlsx([
        [1, "", "", "", "", Q, "Maybe", "?", ""],
        [2, "", "", "", "", "A comment that exists nowhere", "Accept", "ok", ""],
    ])).json()
    assert body["xlsxImport"]["applied"] == []
    reasons = [s["reason"] for s in body["xlsxImport"]["skipped"]]
    assert reasons[0].startswith("no recognized response") and "no matching open comment" in reasons[1]


def test_import_xlsx_skip_summary_groups_identical_reasons(client):
    doc_id = upload(client)
    rows = [[i, "", "", "", "", f"Row {i}", "", "", ""] for i in range(1, 48)]
    body = _upload_xlsx(client, doc_id, _build_crm_xlsx(rows)).json()
    assert len(body["xlsxImport"]["skipped"]) == 47
    assert len(body["xlsxImport"]["skippedSummary"]) == 1
    assert body["xlsxImport"]["skippedSummary"][0]["count"] == 47
    assert len(body["xlsxImport"]["skippedSummary"][0]["rows"]) == 10


def test_import_xlsx_rejects_non_xlsx_and_headerless_sheets(client):
    doc_id = upload(client)
    res = client.post(f"/api/documents/{doc_id}/comments/import-decisions",
                      files={"file": ("matrix.csv", b"a,b", "text/csv")})
    assert res.status_code == 400
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Cover Page"])
    buf = io.BytesIO()
    wb.save(buf)
    assert _upload_xlsx(client, doc_id, buf.getvalue()).status_code == 400


def test_import_creates_a_new_comment_anchored_via_requirement_id(client):
    doc_id = upload(client)
    body = _upload_xlsx(client, doc_id, _build_crm_xlsx([
        [1, "", "", "", "ACA-HRS-0190", "Tighten the verification wording.", "Accept", "Reworded.", ""]])).json()
    applied = body["xlsxImport"]["applied"]
    assert len(applied) == 1 and applied[0]["created"] is True
    c = body["document"]["comments"][applied[0]["comment_id"]]
    assert c["author"] == "CRM Import" and c["shapeId"] == "s1.sh3" and c["quote"] == "ACA-HRS-0190"
    assert c["done"] is True and c["text"] == "Tighten the verification wording."


def test_import_creates_a_new_comment_anchored_on_a_slide_via_ref(client):
    doc_id = upload(client)
    body = _upload_xlsx(client, doc_id, _build_crm_xlsx([
        [1, "", "", "Slide 1", "", "Add the program name.", "Defer", "Next revision.", ""]])).json()
    applied = body["xlsxImport"]["applied"][0]
    c = body["document"]["comments"][applied["comment_id"]]
    assert applied["created"] and c["slideId"] == "s0" and c["shapeId"] == "s0.sh2"  # the slide's title


def test_import_skips_row_when_nothing_locates_a_new_comment(client):
    doc_id = upload(client)
    body = _upload_xlsx(client, doc_id, _build_crm_xlsx([
        [1, "", "", "3.2.2", "ZZZ-9999", "Nowhere to put this.", "Accept", "ok", ""]])).json()
    assert body["xlsxImport"]["applied"] == []
    assert "no Ref #/Para # location found" in body["xlsxImport"]["skipped"][0]["reason"]


# The seven example rows in the real NUWC-issued sample matrix
# (fixtures/crm_sample_from_user.xlsx) are "Test Comment #1".."#7" -- anchor one
# comment per slide-level/shape target with exactly that text so the real file
# can be used as-is (or lightly filled in) against this app's matching logic.
_REAL_SAMPLE_TARGETS = [
    ("s0", None), ("s0", "s0.sh2"), ("s1", "s1.sh3"), ("s2", "s2.sh3"), ("s2", None), ("s3", None), ("s3", "s3.sh3"),
]


def _add_comments_matching_real_sample(client, doc_id: str) -> None:
    for i, (slide, shape_id) in enumerate(_REAL_SAMPLE_TARGETS, start=1):
        res = client.post(f"/api/documents/{doc_id}/comments", json={
            "slide_id": slide, "shape_id": shape_id, "text": f"Test Comment #{i}", "author": "Jane Reviewer"})
        assert res.status_code == 200, res.text


def _fill_real_sample_xlsx(row_updates: dict[int, tuple[str, str]]) -> bytes:
    import openpyxl

    wb = openpyxl.load_workbook(CRM_SAMPLE_XLSX)
    ws = wb["Comments"]
    for row, (response, rationale) in row_updates.items():
        ws.cell(row=row, column=7, value=response)
        ws.cell(row=row, column=8, value=rationale)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_real_sample_xlsx_header_and_rows_are_recognized(client):
    """Ground truth against an actual reviewer-issued file, not a synthetic
    stand-in: 7 unfilled example rows are cleanly reported as "nothing to
    import yet" -- not an error, and not one skip line per row."""
    doc_id = upload(client)
    _add_comments_matching_real_sample(client, doc_id)
    res = _upload_xlsx(client, doc_id, CRM_SAMPLE_XLSX.read_bytes())
    assert res.status_code == 200, res.text
    imp = res.json()["xlsxImport"]
    assert imp["applied"] == [] and len(imp["skipped"]) == 7
    assert [(g["count"], g["reason"]) for g in imp["skippedSummary"]] == [(7, 'no recognized response ("")')]


def test_real_sample_xlsx_round_trips_when_partially_filled_in(client):
    doc_id = upload(client)
    _add_comments_matching_real_sample(client, doc_id)
    xlsx = _fill_real_sample_xlsx({3: ("Accept", "Fixed on slide 1."), 5: ("Reject", "Not applicable.")})
    body = _upload_xlsx(client, doc_id, xlsx).json()
    imp = body["xlsxImport"]
    assert [a["decision"] for a in imp["applied"]] == ["accept", "reject"]
    assert len(imp["skipped"]) == 5
    done = [c for c in body["document"]["comments"].values() if c["done"] and c["author"] == "Jane Reviewer"]
    assert len(done) == 2
